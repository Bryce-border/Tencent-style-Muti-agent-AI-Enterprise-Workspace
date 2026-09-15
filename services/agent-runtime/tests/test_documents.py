import io
import unittest
import zipfile
from app.documents import parse_document, DocumentFiles, MAX_BYTES
from app.config import Settings


class DocumentTests(unittest.TestCase):
    def test_text_encodings_and_csv(self):
        self.assertEqual(parse_document('policy.txt','报销须在30天内提交。'.encode('utf-8-sig')),'报销须在30天内提交。')
        self.assertEqual(parse_document('policy.md','年假申请'.encode('gb18030')),'年假申请')
        self.assertEqual(parse_document('data.csv',b'employee,total\nAlice,30'),'employee | total\nAlice | 30')

    def test_invalid_inputs(self):
        for filename,content in [('bad.exe',b'123'),('empty.txt',b''),('big.txt',b'a'*(MAX_BYTES+1)),('binary.txt',b'abc\0def'),('long.md',b'a'*40001)]:
            with self.subTest(filename=filename),self.assertRaises(ValueError):parse_document(filename,content)

    def test_docx_and_xlsx(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph('采购审批5000元');table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='主管';table.cell(0,1).text='审批'
        stream=io.BytesIO();doc.save(stream)
        self.assertIn('主管 | 审批',parse_document('policy.docx',stream.getvalue()))
        book=Workbook();book.active.append(['员工','报销']);book.active.append(['小明',600]);stream=io.BytesIO();book.save(stream);book.close()
        self.assertIn('小明 | 600',parse_document('data.xlsx',stream.getvalue()))

    def test_empty_pdf_and_zip_bomb(self):
        from pypdf import PdfWriter
        writer=PdfWriter();writer.add_blank_page(width=100,height=100);stream=io.BytesIO();writer.write(stream)
        with self.assertRaisesRegex(ValueError,'OCR'):parse_document('scan.pdf',stream.getvalue())
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w',compression=zipfile.ZIP_DEFLATED) as archive:archive.writestr('test',b'a'*(31*1024*1024))
        with self.assertRaisesRegex(ValueError,'解压'):parse_document('bomb.docx',stream.getvalue())

    def test_object_keys_cannot_cross_workspace_or_document(self):
        files=DocumentFiles(Settings(_env_file=None))
        self.assertEqual(files.key('a','d','a/d/123'),'a/d/123')
        for key in ['b/d/123','a/different/123','a/d/../other']:
            with self.assertRaises(ValueError):files.key('a','d',key)
