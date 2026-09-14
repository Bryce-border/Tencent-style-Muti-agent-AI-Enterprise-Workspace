import asyncio
import base64
import csv
import io
import re
import zipfile
from pathlib import PurePosixPath
from uuid import uuid4

from .config import Settings
from .knowledge import KnowledgeStore

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 40000
EXTENSIONS = {".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx"}


def parse_document(filename: str, content: bytes) -> str:
    extension = PurePosixPath(filename).suffix.lower()
    if extension not in EXTENSIONS or not content or len(content) > MAX_BYTES:
        raise ValueError("仅支持5MB以内的TXT、Markdown、CSV、PDF、DOCX、XLSX文件")
    if extension in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(item.file_size for item in entries) > 30 * 1024 * 1024:
                raise ValueError("文档解压体积超过限制")
    if extension in {".txt", ".md", ".csv"}:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("gb18030")
        if "\0" in text:
            raise ValueError("文本包含无效字符")
        if extension == ".csv":
            rows = csv.reader(io.StringIO(text))
            text = "\n".join(" | ".join(row) for row in rows)
    elif extension == ".pdf":
        from pypdf import PdfReader
        document = PdfReader(io.BytesIO(content))
        if document.is_encrypted or len(document.pages) > 200:
            raise ValueError("不支持加密或超过200页的PDF")
        text = "\n\n".join(f"第{index + 1}页\n{page.extract_text() or ''}" for index, page in enumerate(document.pages))
        if len(re.sub(r"第\d+页|\s", "", text)) < 5:
            raise ValueError("PDF未包含可提取文本，扫描件需要OCR")
    elif extension == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        lines = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            lines.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = "\n".join(lines)
    else:
        from openpyxl import load_workbook
        document = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        lines, cells = [], 0
        try:
            for sheet in document:
                lines.append(f"工作表：{sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > 100000:
                        raise ValueError("工作簿单元格数超过限制")
                    lines.append(" | ".join(str(value) if value is not None else "" for value in row))
        finally:
            document.close()
        text = "\n".join(lines)
    text = text.strip()
    if not text or len(text) > MAX_TEXT:
        raise ValueError("提取文本为空或超过40000字，请拆分文档后重试")
    return text


class DocumentFiles:
    def __init__(self, settings: Settings):
        self.settings = settings

    def client(self):
        from minio import Minio
        return Minio(self.settings.minio_endpoint, access_key=self.settings.minio_access_key,
                     secret_key=self.settings.minio_secret_key, secure=False)

    def key(self, space: str, document: str, key: str):
        if not key.startswith(f"{space}/{document}/") or ".." in key:
            raise ValueError("文件归属不匹配")
        return key

    def store(self, body: dict) -> dict:
        filename = body["filename"]
        data = base64.b64decode(body["content_base64"], validate=True)
        if PurePosixPath(filename).suffix.lower() not in EXTENSIONS or not data or len(data) > MAX_BYTES:
            raise ValueError("仅支持5MB以内的TXT、Markdown、CSV、PDF、DOCX、XLSX文件")
        client = self.client()
        if not client.bucket_exists(self.settings.minio_bucket):
            try:
                client.make_bucket(self.settings.minio_bucket)
            except Exception:
                if not client.bucket_exists(self.settings.minio_bucket):
                    raise
        key = f"{body['workspace_id']}/{body['document_id']}/{uuid4().hex}"
        client.put_object(self.settings.minio_bucket, key, io.BytesIO(data), len(data), content_type="application/octet-stream")
        return {"object_key": key, "byte_size": len(data)}

    def read(self, body: dict) -> bytes:
        key = self.key(body["workspace_id"], body["document_id"], body["object_key"])
        response = self.client().get_object(self.settings.minio_bucket, key)
        try:
            data = response.read(MAX_BYTES + 1)
            if len(data) > MAX_BYTES:
                raise ValueError("文件超过限制")
            return data
        finally:
            response.close()
            response.release_conn()

    async def index(self, body: dict) -> dict:
        data = await asyncio.to_thread(self.read, body)
        # Parse in a disposable subprocess so malformed files have a bounded CPU/runtime cost.
        import multiprocessing
        from concurrent.futures import ProcessPoolExecutor
        executor = ProcessPoolExecutor(max_workers=1, mp_context=multiprocessing.get_context("spawn"))
        try:
            text = await asyncio.wait_for(asyncio.get_running_loop().run_in_executor(executor, parse_document, body["filename"], data), 30)
        finally:
            for process in (getattr(executor, "_processes", None) or {}).values():
                if process.is_alive():
                    process.terminate()
            executor.shutdown(wait=False, cancel_futures=True)
        result = await KnowledgeStore(self.settings).index_document(body["workspace_id"], body["workspace_id"] + "_" + body["document_id"],
            body["title"], text, {"file_document_id": body["document_id"], "version": body["version"], "filename": body["filename"]})
        return {"text": text, "chunk_count": result["chunk_count"]}

    async def exclude(self, body: dict):
        import httpx
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.post(f"{self.settings.elasticsearch_url}/{self.settings.elasticsearch_index}/_update_by_query?refresh=true", json={
                "query": {"bool": {"filter": [{"term": {"workspaceId": body["workspace_id"]}},
                    {"term": {"documentId": body["workspace_id"] + "_" + body["document_id"]}}]}},
                "script": {"source": "ctx._source.metadata.retrieval_excluded = params.excluded", "params": {"excluded": body["excluded"]}}})
            response.raise_for_status()
            if response.json().get("failures"):
                raise RuntimeError("document index update failed")
        return {"ok": True}
